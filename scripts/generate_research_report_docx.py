from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import matplotlib
import mistune
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image


matplotlib.use("Agg")
import matplotlib.pyplot as plt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = PROJECT_ROOT / "reports" / "research_report_draft.md"
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "reports"
    / "final"
    / "基于预定义知识图谱的轻量化混合GraphRAG科研实践报告.docx"
)
DEFAULT_MANIFEST = PROJECT_ROOT / "reports" / "final" / "research_report_docx_manifest.json"

REPORT_TITLE = "基于预定义知识图谱的轻量化混合 GraphRAG 机器学习知识问答 Agent"
REPORT_SUBTITLE = "科研实践报告"
REPORT_DATE = "2026 年 7 月 24 日"
KEYWORDS = "GraphRAG；知识图谱；检索增强生成；LangGraph；证据验证"

# narrative_proposal with the named academic_a4_cn override.
PAGE_WIDTH_TWIPS = 11906
PAGE_HEIGHT_TWIPS = 16838
PAGE_MARGIN_TWIPS = 1440
HEADER_FOOTER_TWIPS = 708
CONTENT_WIDTH_DXA = 9026
TABLE_INDENT_DXA = 120
BODY_SIZE_PT = 11
BODY_LINE_SPACING = 1.3
BODY_AFTER_PT = 1
BODY_FIRST_LINE_PT = 22
BODY_EAST_ASIA_FONT = "SimSun"
HEADING_EAST_ASIA_FONT = "SimHei"
LATIN_FONT = "Times New Roman"
CODE_FONT = "Consolas"
HEADING_COLOR = "1F4E79"
HEADING_DARK_COLOR = "17365D"
MUTED_COLOR = "5B6573"
TABLE_HEADER_FILL = "E8EEF5"
TABLE_BORDER_COLOR = "AAB7C4"
CODE_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
CALLOUT_BORDER = "3C6E71"

# Static pagination is filled after the rendered QA pass. Keeping the values
# in the generator makes the final TOC deterministic in Word and LibreOffice.
STATIC_TOC_PAGE_LABELS: dict[str, str] = {
    "第 1 章 绪论": "1",
    "第 2 章 相关技术": "3",
    "第 3 章 数据与知识库构建": "4",
    "第 4 章 系统总体架构": "6",
    "第 5 章 核心方法": "8",
    "第 6 章 系统实现": "13",
    "第 7 章 实验设计与结果": "15",
    "第 8 章 总结与展望": "25",
    "参考文献": "26",
    "附录 A 可复现实验命令": "27",
}

TABLE_CAPTIONS = [
    "表 3-1  官方文档来源",
    "表 3-2  各来源 Chunk 数量",
    "表 3-3  图谱关系类型与数量",
    "表 3-4  评测数据集划分",
    "表 4-1  LangGraph 节点输入输出",
    "表 6-1  软件环境",
    "表 6-2  主要模块与职责",
    "表 7-1  四种实验方法配置",
    "表 7-2  Final 冻结结果",
    "表 7-3  Pilot 主结果",
    "表 7-4  No Verifier 消融配置与结果",
    "表 7-5  Final 错误案例 T-DF-01",
    "表 7-6  Pilot 典型误差",
    "表 7-7  v2 Extension 用户确认结果",
]

DISPLAY_FORMULAS = {
    "s_i": r"$s_i = \frac{v_q \cdot v_i}{\left|v_q\right|\,\left|v_i\right|}$",
    "S = 0.30": r"$S = 0.30C + 0.25V + 0.20P + 0.25R$",
    "DecisionAccuracy": (
        r"$\mathrm{DecisionAccuracy} = "
        r"\frac{\#\,\mathrm{correct\ pass/refuse}}{N}$"
    ),
    "AnswerCorrectness": r"$\mathrm{AnswerCorrectness} = \frac{\sum_i c_i}{2N}$",
    "Faithfulness": r"$\mathrm{Faithfulness} = \frac{\sum_i f_i}{2N}$",
    "Recall@5": (
        r"$\mathrm{Recall@5} = "
        r"\frac{\#\,\mathrm{questions\ with\ gold\ chunk\ in\ top5}}"
        r"{\#\,\mathrm{eligible\ answerable\ questions}}$"
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def set_attribute(element, name: str, value: str) -> None:
    element.set(qn(name), value)


def set_run_fonts(
    run,
    *,
    east_asia: str = BODY_EAST_ASIA_FONT,
    latin: str = LATIN_FONT,
) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        set_attribute(r_fonts, attribute, latin)
    set_attribute(r_fonts, "w:eastAsia", east_asia)


def set_style_fonts(style, *, east_asia: str, latin: str) -> None:
    style.font.name = latin
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        set_attribute(r_fonts, attribute, latin)
    set_attribute(r_fonts, "w:eastAsia", east_asia)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    set_attribute(shd, "w:fill", fill)
    set_attribute(shd, "w:val", "clear")


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    set_attribute(shd, "w:fill", fill)
    set_attribute(shd, "w:val", "clear")


def set_paragraph_left_border(paragraph, color: str, size: int = 16) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    set_attribute(left, "w:val", "single")
    set_attribute(left, "w:sz", str(size))
    set_attribute(left, "w:space", "8")
    set_attribute(left, "w:color", color)


def set_cell_margins(cell, *, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        set_attribute(node, "w:w", str(value))
        set_attribute(node, "w:type", "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    set_attribute(marker, "w:val", "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:cantSplit"))
    if marker is None:
        marker = OxmlElement("w:cantSplit")
        tr_pr.append(marker)


def configure_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        set_attribute(node, "w:val", "single")
        set_attribute(node, "w:sz", "5")
        set_attribute(node, "w:space", "0")
        set_attribute(node, "w:color", TABLE_BORDER_COLOR)


def configure_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError("table column widths do not sum to the content width")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    set_attribute(tbl_w, "w:w", str(CONTENT_WIDTH_DXA))
    set_attribute(tbl_w, "w:type", "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    set_attribute(tbl_ind, "w:w", str(TABLE_INDENT_DXA))
    set_attribute(tbl_ind, "w:type", "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    set_attribute(layout, "w:type", "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        set_attribute(grid_col, "w:w", str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            set_attribute(tc_w, "w:w", str(width))
            set_attribute(tc_w, "w:type", "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=70, bottom=70, start=105, end=105)
    configure_table_borders(table)


def display_width(text: str) -> float:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in text)


def calculate_column_widths(rows: list[list[str]]) -> list[int]:
    column_count = len(rows[0])
    max_lengths = [
        max(display_width(row[index]) for row in rows)
        for index in range(column_count)
    ]
    weights = [max(5.0, min(math.sqrt(length + 4) * 3.0, 18.0)) for length in max_lengths]

    headers = rows[0]
    if headers and headers[0] in {"ID", "指标", "项目", "组件", "节点", "关系", "来源"}:
        weights[0] = min(weights[0], 8.0)
    if headers and headers[0] == "方法":
        weights[0] = max(weights[0], 14.0)

    minimum = 650 if column_count >= 7 else 800 if column_count >= 5 else 1000
    raw = [CONTENT_WIDTH_DXA * weight / sum(weights) for weight in weights]
    widths = [max(minimum, int(value)) for value in raw]

    while sum(widths) > CONTENT_WIDTH_DXA:
        candidates = [index for index, value in enumerate(widths) if value > minimum]
        if not candidates:
            break
        largest = max(candidates, key=lambda index: widths[index] - minimum)
        widths[largest] -= min(10, sum(widths) - CONTENT_WIDTH_DXA)
    if sum(widths) < CONTENT_WIDTH_DXA:
        widest = max(range(column_count), key=lambda index: max_lengths[index])
        widths[widest] += CONTENT_WIDTH_DXA - sum(widths)
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError("unable to allocate exact table geometry")
    return widths


def add_field(run, instruction: str, display: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    set_attribute(begin, "w:fldCharType", "begin")
    instr = OxmlElement("w:instrText")
    set_attribute(instr, "xml:space", "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    set_attribute(separate, "w:fldCharType", "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    set_attribute(end, "w:fldCharType", "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def set_page_number_format(section, *, style: str, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    set_attribute(pg_num, "w:fmt", style)
    set_attribute(pg_num, "w:start", str(start))


def configure_section(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def clear_story(story) -> None:
    for paragraph in story.paragraphs:
        for run in list(paragraph.runs):
            paragraph._p.remove(run._r)


def configure_header_footer(section, *, front_matter: bool = False) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_story(section.header)
    clear_story(section.footer)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("科研实践报告" if front_matter else "轻量化混合 GraphRAG 机器学习知识问答 Agent")
    set_run_fonts(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED_COLOR)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    left = footer.add_run("- ")
    set_run_fonts(left)
    page = footer.add_run()
    set_run_fonts(page)
    add_field(page, "PAGE", "1")
    right = footer.add_run(" -")
    set_run_fonts(right)
    for item in (left, page, right):
        item.font.size = Pt(9)
        item.font.color.rgb = RGBColor.from_string(MUTED_COLOR)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    set_style_fonts(normal, east_asia=BODY_EAST_ASIA_FONT, latin=LATIN_FONT)
    normal.font.size = Pt(BODY_SIZE_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER_PT)
    normal.paragraph_format.line_spacing = BODY_LINE_SPACING
    normal.paragraph_format.first_line_indent = Pt(BODY_FIRST_LINE_PT)

    title = styles["Title"]
    set_style_fonts(title, east_asia=HEADING_EAST_ASIA_FONT, latin=LATIN_FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(HEADING_DARK_COLOR)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.2

    subtitle = styles["Subtitle"]
    set_style_fonts(subtitle, east_asia=HEADING_EAST_ASIA_FONT, latin=LATIN_FONT)
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(HEADING_COLOR)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)

    heading_specs = {
        "Heading 1": (16, HEADING_DARK_COLOR, 0, 16, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (14, HEADING_COLOR, 10, 5, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": (12, HEADING_DARK_COLOR, 8, 3, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for name, (size, color, before, after, alignment) in heading_specs.items():
        style = styles[name]
        set_style_fonts(style, east_asia=HEADING_EAST_ASIA_FONT, latin=LATIN_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.first_line_indent = Pt(0)

    styles["Heading 1"].paragraph_format.page_break_before = True

    caption = styles["Caption"]
    set_style_fonts(caption, east_asia=BODY_EAST_ASIA_FONT, latin=LATIN_FONT)
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(MUTED_COLOR)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_style_fonts(style, east_asia=BODY_EAST_ASIA_FONT, latin=LATIN_FONT)
        style.font.size = Pt(BODY_SIZE_PT)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.first_line_indent = Pt(0)

    if "Code Block" not in styles:
        styles.add_style("Code Block", 1)
    code = styles["Code Block"]
    set_style_fonts(code, east_asia="Microsoft YaHei", latin=CODE_FONT)
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("243B53")
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.first_line_indent = Pt(0)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.15
    code.paragraph_format.keep_together = True

    if "Reference" not in styles:
        styles.add_style("Reference", 1)
    reference = styles["Reference"]
    set_style_fonts(reference, east_asia=BODY_EAST_ASIA_FONT, latin=LATIN_FONT)
    reference.font.size = Pt(10)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.left_indent = Inches(0.28)
    reference.paragraph_format.first_line_indent = Inches(-0.28)
    reference.paragraph_format.space_before = Pt(0)
    reference.paragraph_format.space_after = Pt(5)
    reference.paragraph_format.line_spacing = 1.25


def create_numbering(document: Document) -> dict[str, int]:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    next_abstract = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1
    result: dict[str, int] = {"next_num_id": next_num}

    for index, (name, fmt, marker) in enumerate(
        (("bullet", "bullet", "•"), ("number", "decimal", "%1."))
    ):
        abstract_id = next_abstract + index
        abstract = OxmlElement("w:abstractNum")
        set_attribute(abstract, "w:abstractNumId", str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        set_attribute(multi, "w:val", "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        set_attribute(level, "w:ilvl", "0")
        start = OxmlElement("w:start")
        set_attribute(start, "w:val", "1")
        num_fmt = OxmlElement("w:numFmt")
        set_attribute(num_fmt, "w:val", fmt)
        lvl_text = OxmlElement("w:lvlText")
        set_attribute(lvl_text, "w:val", marker)
        lvl_jc = OxmlElement("w:lvlJc")
        set_attribute(lvl_jc, "w:val", "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        set_attribute(tab, "w:val", "num")
        set_attribute(tab, "w:pos", "540")
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        set_attribute(indent, "w:left", "540")
        set_attribute(indent, "w:hanging", "280")
        spacing = OxmlElement("w:spacing")
        set_attribute(spacing, "w:after", "80")
        set_attribute(spacing, "w:line", "300")
        set_attribute(spacing, "w:lineRule", "auto")
        p_pr.extend((tabs, indent, spacing))
        level.extend((start, num_fmt, lvl_text, lvl_jc, p_pr))
        abstract.append(level)
        numbering.append(abstract)
        result[f"{name}_abstract_id"] = abstract_id
    return result


def create_list_instance(numbering_ids: dict[str, int], list_type: str, document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = numbering_ids[f"{list_type}_abstract_id"]
    num_id = numbering_ids["next_num_id"]
    numbering_ids["next_num_id"] += 1

    num = OxmlElement("w:num")
    set_attribute(num, "w:numId", str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    set_attribute(abstract_ref, "w:val", str(abstract_id))
    num.append(abstract_ref)
    if list_type == "number":
        override = OxmlElement("w:lvlOverride")
        set_attribute(override, "w:ilvl", "0")
        start = OxmlElement("w:startOverride")
        set_attribute(start, "w:val", "1")
        override.append(start)
        num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    set_attribute(ilvl, "w:val", "0")
    num = OxmlElement("w:numId")
    set_attribute(num, "w:val", str(num_id))
    num_pr.extend((ilvl, num))


def flatten_inline(nodes: list[dict]) -> str:
    pieces: list[str] = []
    for node in nodes:
        if "text" in node:
            pieces.append(node["text"])
        if "children" in node:
            pieces.append(flatten_inline(node["children"]))
    return "".join(pieces)


def clean_inline_math(text: str) -> str:
    replacements = {
        r"\in": "∈",
        r"\{": "{",
        r"\}": "}",
        r"\cdot": "·",
        r"\sum": "Σ",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.replace("$", "")


def requires_left_alignment(text: str) -> bool:
    """Avoid stretched spacing in long hashes, paths, and URLs."""
    return bool(
        "SHA-256" in text
        or re.search(r"\b[a-f0-9]{32,}\b", text, flags=re.IGNORECASE)
        or (len(text) > 100 and ("http://" in text or "https://" in text))
        or bool(re.search(r"\b(?:reports|scripts|data|src|config)/[\w./-]+", text))
    )


def add_text_with_inline_math(paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    parts = re.split(r"(\$[^$]+\$)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(clean_inline_math(part))
        set_run_fonts(run)
        run.bold = bold
        run.italic = italic or (part.startswith("$") and part.endswith("$"))
        if part.startswith("$"):
            run.font.name = "Cambria Math"
            set_run_fonts(run, east_asia="Cambria Math", latin="Cambria Math")


def add_inline_nodes(paragraph, nodes: list[dict], *, bold: bool = False, italic: bool = False) -> None:
    for node in nodes:
        node_type = node["type"]
        if node_type == "text":
            add_text_with_inline_math(paragraph, node["text"], bold=bold, italic=italic)
        elif node_type == "codespan":
            run = paragraph.add_run(node["text"])
            set_run_fonts(run, east_asia="Microsoft YaHei", latin=CODE_FONT)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string("8B3A3A")
            set_run_shading(run, "F3F4F6")
        elif node_type == "strong":
            add_inline_nodes(paragraph, node.get("children", []), bold=True, italic=italic)
        elif node_type == "emphasis":
            add_inline_nodes(paragraph, node.get("children", []), bold=bold, italic=True)
        elif node_type == "image":
            continue
        elif "children" in node:
            add_inline_nodes(paragraph, node["children"], bold=bold, italic=italic)


def set_run_shading(run, fill: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    shd = r_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        r_pr.append(shd)
    set_attribute(shd, "w:fill", fill)
    set_attribute(shd, "w:val", "clear")


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.LINE)
    paragraph.add_run().add_break(WD_BREAK.LINE)
    paragraph.add_run().add_break(WD_BREAK.LINE)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("科 研 实 践 报 告")
    set_run_fonts(run, east_asia=HEADING_EAST_ASIA_FONT)
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("3C6E71")

    title = document.add_paragraph(style="Title")
    title.paragraph_format.keep_together = True
    title.add_run(REPORT_TITLE)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("scikit-learn 知识问答 Agent")

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(82)

    metadata = [
        ("项目仓库", "zzh1126/RAGagent"),
        ("基线版本", "v1.0-baseline"),
        ("扩展实验", "extension-qwen3-4b-v2-e207cb91"),
        ("完成日期", REPORT_DATE),
    ]
    for label, value in metadata:
        line = document.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_after = Pt(6)
        line.paragraph_format.first_line_indent = Pt(0)
        label_run = line.add_run(f"{label}：")
        set_run_fonts(label_run, east_asia=HEADING_EAST_ASIA_FONT)
        label_run.bold = True
        label_run.font.size = Pt(10.5)
        value_run = line.add_run(value)
        set_run_fonts(value_run)
        value_run.font.size = Pt(10.5)


def add_abstract(document: Document, abstract_nodes: list[dict]) -> None:
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = False
    heading.add_run("摘要")
    keywords_text = KEYWORDS
    for node in abstract_nodes:
        if node["type"] != "paragraph":
            continue
        raw = flatten_inline(node.get("children", [])).strip()
        if raw.startswith("关键词："):
            keywords_text = raw.removeprefix("关键词：").strip()
            continue
        paragraph = document.add_paragraph()
        add_inline_nodes(paragraph, node["children"])
    keywords = document.add_paragraph()
    keywords.paragraph_format.first_line_indent = Pt(0)
    keywords.paragraph_format.space_before = Pt(8)
    label = keywords.add_run("关键词：")
    set_run_fonts(label, east_asia=HEADING_EAST_ASIA_FONT)
    label.bold = True
    add_text_with_inline_math(keywords, keywords_text)


def normalize_toc_heading(text: str) -> str:
    if text == "参考文献（草稿）":
        return "参考文献"
    return text


def collect_toc_entries(nodes: list[dict], body_start: int) -> list[tuple[str, str]]:
    entries = [("摘要", "i")]
    for node in nodes[body_start:]:
        if node["type"] != "heading" or node["level"] != 1:
            continue
        text = normalize_toc_heading(flatten_inline(node["children"]))
        if text.startswith("附录 B"):
            break
        entries.append((text, STATIC_TOC_PAGE_LABELS.get(text, "—")))
    return entries


def add_toc(document: Document, entries: list[tuple[str, str]]) -> None:
    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = False
    heading.add_run("目录")
    for text, page_label in entries:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.08), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        title_run = paragraph.add_run(text)
        set_run_fonts(title_run)
        title_run.font.size = Pt(10.5)
        page_run = paragraph.add_run(f"\t{page_label}")
        set_run_fonts(page_run)
        page_run.font.size = Pt(10.5)


def formula_image(latex: str) -> io.BytesIO:
    fig = plt.figure(figsize=(8.0, 0.65), dpi=180)
    fig.patch.set_alpha(0)
    fig.text(0.5, 0.5, latex, ha="center", va="center", fontsize=15, color="#102A43")
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", transparent=True, bbox_inches="tight", pad_inches=0.06)
    plt.close(fig)
    buffer.seek(0)
    return buffer


def add_display_formula(document: Document, raw: str) -> None:
    latex = None
    for marker, value in DISPLAY_FORMULAS.items():
        if marker in raw:
            latex = value
            break
    if latex is None:
        latex = f"${clean_inline_math(raw.replace('$$', '').strip())}$"
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run()
    wide_formula = "Recall@5" in latex or len(latex) > 90
    run.add_picture(formula_image(latex), width=Inches(5.5 if wide_formula else 3.8))


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    set_paragraph_shading(paragraph, CODE_FILL)
    set_paragraph_left_border(paragraph, "8093A7", size=10)
    for index, line in enumerate(text.rstrip().splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_fonts(run, east_asia="Microsoft YaHei", latin=CODE_FONT)
        run.font.size = Pt(8.5)


def add_callout(document: Document, node: dict) -> None:
    text = flatten_inline(node.get("children", []))
    if text.startswith("版本："):
        return
    if text.startswith("图 4-1"):
        text = "来源：静态架构图由 scripts/generate_report_figures.py 根据当前实现生成。"
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.25
    if requires_left_alignment(text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_shading(paragraph, CALLOUT_FILL)
    set_paragraph_left_border(paragraph, CALLOUT_BORDER)
    add_text_with_inline_math(paragraph, text)


def add_image(document: Document, image_node: dict) -> None:
    path = SOURCE_PATH.parent / image_node["src"]
    if not path.is_file():
        raise FileNotFoundError(f"missing report figure: {path}")
    with Image.open(path) as image:
        pixel_width, pixel_height = image.size
    aspect = pixel_width / pixel_height
    max_width = 6.05
    max_height = 5.25
    width = min(max_width, max_height * aspect)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", image_node["alt"])

    caption = document.add_paragraph(style="Caption")
    caption.paragraph_format.keep_with_next = False
    caption.add_run(image_node["alt"])


def table_rows(node: dict) -> list[list[dict]]:
    header = node["children"][0]
    body = node["children"][1]
    rows = [[cell for cell in header["children"]]]
    rows.extend([[cell for cell in row["children"]] for row in body["children"]])
    return rows


def add_table(document: Document, node: dict, table_index: int) -> None:
    cells = table_rows(node)
    plain_rows = [[flatten_inline(cell["children"]) for cell in row] for row in cells]
    widths = calculate_column_widths(plain_rows)
    column_count = len(cells[0])

    caption = document.add_paragraph(style="Caption")
    caption.paragraph_format.keep_with_next = True
    caption.add_run(TABLE_CAPTIONS[table_index])

    table = document.add_table(rows=len(cells), cols=column_count)
    configure_table_geometry(table, widths)
    font_size = 7.5 if column_count >= 8 else 8 if column_count >= 6 else 8.8 if column_count >= 4 else 9.2
    for row_index, row in enumerate(cells):
        table_row = table.rows[row_index]
        prevent_row_split(table_row)
        if row_index == 0:
            set_repeat_table_header(table_row)
        for column_index, cell_node in enumerate(row):
            cell = table_row.cells[column_index]
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            text = plain_rows[row_index][column_index]
            numeric = bool(re.fullmatch(r"[\d./%()\-+\s]+", text))
            short = display_width(text) <= 15
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if row_index == 0 or numeric or (short and column_index == 0)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_nodes(paragraph, cell_node["children"])
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
                    set_run_fonts(run, east_asia=HEADING_EAST_ASIA_FONT)


def add_list(document: Document, node: dict, numbering: dict[str, int]) -> None:
    list_type = "number" if node.get("ordered") else "bullet"
    style = "List Number" if list_type == "number" else "List Bullet"
    num_id = create_list_instance(numbering, list_type, document)
    for item in node["children"]:
        paragraph = document.add_paragraph(style=style)
        apply_numbering(paragraph, num_id)
        paragraph.paragraph_format.keep_together = True
        children = item.get("children", [])
        inline_nodes: list[dict] = []
        for child in children:
            if child["type"] in {"block_text", "paragraph"}:
                inline_nodes.extend(child.get("children", []))
        add_inline_nodes(paragraph, inline_nodes)


def add_body_paragraph(document: Document, node: dict, *, references: bool) -> None:
    raw = flatten_inline(node.get("children", []))
    if raw.strip().startswith("$$") and raw.strip().endswith("$$"):
        add_display_formula(document, raw)
        return
    image_nodes = [child for child in node.get("children", []) if child["type"] == "image"]
    if image_nodes:
        for image_node in image_nodes:
            add_image(document, image_node)
        return
    style = "Reference" if references else None
    paragraph = document.add_paragraph(style=style)
    if not references and requires_left_alignment(raw):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
    add_inline_nodes(paragraph, node.get("children", []))


def parse_source() -> list[dict]:
    markdown = mistune.create_markdown(renderer="ast", plugins=["table", "strikethrough"])
    return markdown(SOURCE_PATH.read_text(encoding="utf-8"))


def split_abstract(nodes: list[dict]) -> tuple[list[dict], int]:
    abstract_start = next(
        index
        for index, node in enumerate(nodes)
        if node["type"] == "heading"
        and node["level"] == 2
        and flatten_inline(node["children"]) == "摘要"
    )
    body_start = next(
        index
        for index, node in enumerate(nodes)
        if node["type"] == "heading"
        and node["level"] == 1
        and flatten_inline(node["children"]).startswith("第 1 章")
    )
    return nodes[abstract_start + 1 : body_start], body_start


def set_document_settings(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    set_attribute(update, "w:val", "true")

    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    setting = OxmlElement("w:compatSetting")
    set_attribute(setting, "w:name", "compatibilityMode")
    set_attribute(setting, "w:uri", "http://schemas.microsoft.com/office/word")
    set_attribute(setting, "w:val", "15")
    compat.append(setting)


def build_document(nodes: list[dict]) -> tuple[Document, dict]:
    document = Document()
    configure_styles(document)
    set_document_settings(document)
    numbering = create_numbering(document)
    for section in document.sections:
        configure_section(section)
    cover_section = document.sections[0]
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    clear_story(cover_section.header)
    clear_story(cover_section.footer)

    properties = document.core_properties
    properties.title = REPORT_TITLE
    properties.subject = "scikit-learn 知识问答 Agent 科研实践报告"
    properties.author = "RAGagent Project"
    properties.keywords = KEYWORDS
    properties.comments = "Generated from the validated Markdown research report."
    properties.created = datetime(2026, 7, 24, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 7, 24, tzinfo=timezone.utc)

    add_cover(document)
    front_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front_section)
    configure_header_footer(front_section, front_matter=True)
    set_page_number_format(front_section, style="lowerRoman", start=1)

    abstract_nodes, body_start = split_abstract(nodes)
    toc_entries = collect_toc_entries(nodes, body_start)
    add_abstract(document, abstract_nodes)
    add_toc(document, toc_entries)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    configure_header_footer(body_section, front_matter=False)
    set_page_number_format(body_section, style="decimal", start=1)

    table_index = 0
    references = False
    appendix_b_skipped = False
    heading_count = 0
    paragraph_count = 0
    for node in nodes[body_start:]:
        node_type = node["type"]
        if node_type == "heading":
            text = flatten_inline(node["children"])
            if text.startswith("附录 B"):
                appendix_b_skipped = True
                break
            if text == "参考文献（草稿）":
                text = "参考文献"
            references = text == "参考文献" or references and not text.startswith("附录 A")
            if text.startswith("附录 A"):
                references = False
            paragraph = document.add_paragraph(style=f"Heading {node['level']}")
            paragraph.add_run(text)
            heading_count += 1
        elif node_type == "paragraph":
            add_body_paragraph(document, node, references=references)
            paragraph_count += 1
        elif node_type == "list":
            add_list(document, node, numbering)
        elif node_type == "table":
            add_table(document, node, table_index)
            table_index += 1
        elif node_type == "block_code":
            add_code_block(document, node["text"])
        elif node_type == "block_quote":
            add_callout(document, node)
        elif node_type in {"thematic_break", "newline"}:
            continue

    if table_index != len(TABLE_CAPTIONS):
        raise ValueError(f"expected {len(TABLE_CAPTIONS)} tables, generated {table_index}")

    stats = {
        "source_ast_node_count": len(nodes),
        "heading_count": heading_count + 2,
        "body_paragraph_node_count": paragraph_count,
        "table_count": table_index,
        "figure_count": 9,
        "section_count": len(document.sections),
        "toc_entries": toc_entries,
        "appendix_b_internal_checklist_omitted": appendix_b_skipped,
    }
    return document, stats


def figure_sources(nodes: list[dict]) -> list[Path]:
    paths: list[Path] = []
    for node in nodes:
        if node["type"] != "paragraph":
            continue
        for child in node.get("children", []):
            if child["type"] == "image":
                paths.append(SOURCE_PATH.parent / child["src"])
    if len(paths) != 9 or len(set(paths)) != 9:
        raise ValueError("research report must reference exactly nine unique figures")
    return paths


def build_manifest(output: Path, nodes: list[dict], stats: dict) -> dict:
    figures = figure_sources(nodes)
    return {
        "schema_version": "1.1",
        "artifact": "research_report_final_docx",
        "title": REPORT_TITLE,
        "source": {
            "path": str(SOURCE_PATH.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "sha256": sha256(SOURCE_PATH),
        },
        "generator": {
            "path": "scripts/generate_research_report_docx.py",
            "sha256": sha256(Path(__file__)),
        },
        "figures": {
            str(path.relative_to(PROJECT_ROOT)).replace("\\", "/"): sha256(path)
            for path in figures
        },
        "toc": {
            "mode": "static_render_verified",
            "entries": [
                {"heading": heading, "page_label": page_label}
                for heading, page_label in stats["toc_entries"]
            ],
        },
        "output": {
            "path": str(output.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "sha256": sha256(output),
            "size_bytes": output.stat().st_size,
        },
        "design": {
            "preset": "narrative_proposal",
            "named_override": "academic_a4_cn",
            "page_size_twips": [PAGE_WIDTH_TWIPS, PAGE_HEIGHT_TWIPS],
            "margins_twips": [PAGE_MARGIN_TWIPS] * 4,
            "header_footer_twips": HEADER_FOOTER_TWIPS,
            "content_width_dxa": CONTENT_WIDTH_DXA,
            "body_font_east_asia": BODY_EAST_ASIA_FONT,
            "body_font_latin": LATIN_FONT,
            "body_size_pt": BODY_SIZE_PT,
            "body_line_spacing": BODY_LINE_SPACING,
            "body_after_pt": BODY_AFTER_PT,
            "body_first_line_pt": BODY_FIRST_LINE_PT,
            "table_indent_dxa": TABLE_INDENT_DXA,
            "table_header_fill": TABLE_HEADER_FILL,
        },
        "structure": stats,
        "limitations": [
            "The cover omits unknown student and institution metadata rather than inventing placeholders.",
            "Appendix B is an internal delivery checklist and is intentionally omitted from the formal report.",
            "The DOCX uses a static, render-verified chapter table of contents and updateable page-number fields.",
        ],
    }


def generate(output: Path, manifest_path: Path) -> None:
    nodes = parse_source()
    document, stats = build_document(nodes)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    manifest = build_manifest(output, nodes, stats)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(
        f"OK: generated final report DOCX tables={stats['table_count']} "
        f"figures={stats['figure_count']} output={output}"
    )


def read_docx_xml(output: Path) -> tuple[bytes, bytes]:
    with ZipFile(output, "r") as archive:
        return archive.read("word/document.xml"), archive.read("word/styles.xml")


def validate_table_geometry(document_xml: bytes) -> list[str]:
    from xml.etree import ElementTree as ET

    root = ET.fromstring(document_xml)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    errors: list[str] = []
    tables = root.findall(".//w:tbl", namespace)
    if len(tables) != len(TABLE_CAPTIONS):
        errors.append(f"expected {len(TABLE_CAPTIONS)} tables, found {len(tables)}")
    for index, table in enumerate(tables, start=1):
        width = table.find("w:tblPr/w:tblW", namespace)
        indent = table.find("w:tblPr/w:tblInd", namespace)
        layout = table.find("w:tblPr/w:tblLayout", namespace)
        grid = table.findall("w:tblGrid/w:gridCol", namespace)
        grid_widths = [int(column.get(qn("w:w"), "0")) for column in grid]
        if width is None or width.get(qn("w:w")) != str(CONTENT_WIDTH_DXA):
            errors.append(f"table {index}: tblW is not {CONTENT_WIDTH_DXA}")
        if indent is None or indent.get(qn("w:w")) != str(TABLE_INDENT_DXA):
            errors.append(f"table {index}: tblInd is not {TABLE_INDENT_DXA}")
        if layout is None or layout.get(qn("w:type")) != "fixed":
            errors.append(f"table {index}: fixed layout is missing")
        if sum(grid_widths) != CONTENT_WIDTH_DXA:
            errors.append(f"table {index}: grid widths do not sum to {CONTENT_WIDTH_DXA}")
        for row_number, row in enumerate(table.findall("w:tr", namespace), start=1):
            cell_widths = [
                int(cell.find("w:tcPr/w:tcW", namespace).get(qn("w:w"), "0"))
                for cell in row.findall("w:tc", namespace)
            ]
            if cell_widths != grid_widths:
                errors.append(f"table {index} row {row_number}: tcW differs from tblGrid")
    return errors


def check(output: Path, manifest_path: Path) -> None:
    if not output.is_file() or not manifest_path.is_file():
        raise FileNotFoundError("final DOCX or manifest is missing")
    nodes = parse_source()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    _, body_start = split_abstract(nodes)
    expected_toc_entries = [
        {"heading": heading, "page_label": page_label}
        for heading, page_label in collect_toc_entries(nodes, body_start)
    ]
    if any(entry["page_label"] == "—" for entry in expected_toc_entries):
        raise ValueError("static TOC page labels have not been populated after render QA")
    expected_sources = {
        str(path.relative_to(PROJECT_ROOT)).replace("\\", "/"): sha256(path)
        for path in figure_sources(nodes)
    }
    checks = {
        "artifact": "research_report_final_docx",
        "source.path": "reports/research_report_draft.md",
        "source.sha256": sha256(SOURCE_PATH),
        "generator.sha256": sha256(Path(__file__)),
        "output.sha256": sha256(output),
        "figures": expected_sources,
        "toc.entries": expected_toc_entries,
    }
    actuals = {
        "artifact": manifest.get("artifact"),
        "source.path": manifest.get("source", {}).get("path"),
        "source.sha256": manifest.get("source", {}).get("sha256"),
        "generator.sha256": manifest.get("generator", {}).get("sha256"),
        "output.sha256": manifest.get("output", {}).get("sha256"),
        "figures": manifest.get("figures"),
        "toc.entries": manifest.get("toc", {}).get("entries"),
    }
    for field, expected in checks.items():
        if actuals[field] != expected:
            raise ValueError(f"final DOCX manifest field is stale: {field}")

    document = Document(output)
    if len(document.sections) != 3:
        raise ValueError(f"expected three document sections, found {len(document.sections)}")
    for index, section in enumerate(document.sections, start=1):
        if abs(section.page_width.twips - PAGE_WIDTH_TWIPS) > 2:
            raise ValueError(f"section {index}: page width is not A4")
        if abs(section.page_height.twips - PAGE_HEIGHT_TWIPS) > 2:
            raise ValueError(f"section {index}: page height is not A4")
        margins = (
            section.top_margin.twips,
            section.right_margin.twips,
            section.bottom_margin.twips,
            section.left_margin.twips,
        )
        if any(abs(value - PAGE_MARGIN_TWIPS) > 2 for value in margins):
            raise ValueError(f"section {index}: margins differ from the design tokens")

    document_xml, _ = read_docx_xml(output)
    xml_text = document_xml.decode("utf-8")
    for forbidden in (
        "[[TOC]]",
        "请在 Word 中更新目录",
        "科研实践报告草稿",
        "附录 B 待完成清单",
    ):
        if forbidden in xml_text:
            raise ValueError(f"final DOCX contains forbidden draft text: {forbidden}")
    if 'TOC \\o' in xml_text:
        raise ValueError("final DOCX still contains a dynamic TOC field")
    if xml_text.count("<wp:inline") != 15:
        raise ValueError("expected 9 figures plus 6 display-formula images")
    geometry_errors = validate_table_geometry(document_xml)
    if geometry_errors:
        raise ValueError("; ".join(geometry_errors[:5]))
    print(
        "OK: final report DOCX and manifest are current "
        f"sections=3 tables={len(TABLE_CAPTIONS)} figures=9 formulas=6"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate or validate the final research-report DOCX."
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--check", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output = args.output if args.output.is_absolute() else PROJECT_ROOT / args.output
    manifest = args.manifest if args.manifest.is_absolute() else PROJECT_ROOT / args.manifest
    try:
        if args.check:
            check(output, manifest)
        else:
            generate(output, manifest)
    except (FileNotFoundError, ValueError, OSError, KeyError, IndexError) as exc:
        print(f"ERROR: {exc}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
